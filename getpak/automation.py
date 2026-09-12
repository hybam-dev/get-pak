import os
import sys
import ast
import json
import inspect
import hashlib
import platform
import time
import re
from importlib import metadata as package_metadata
import numpy as np
import pandas as pd
import xarray as xr

from pathlib import Path
from datetime import datetime, timezone
from getpak import inversion_functions as ifunc
from getpak.input import Input
from getpak.input import GRS as g
from getpak.input import ACOLITE_S2
from getpak.output import Raster as r
from getpak.commons import Utils as u
from getpak.methods import Methods

grs = g()
i = Input()
m = Methods()

class Pipelines:
    
    def __init__(self, config_path=None):
        # GET-Pak settings
        self.settings = u.resolve_encoding_settings(u.read_config(config_path=config_path))
        self.encoding_settings = self.settings
        self._skip_output_targets = set()
        self.grid_by_uid = {}
        self.settings['_config_dir'] = str(Path(config_path).resolve().parent) if config_path else str(Path(__file__).resolve().parent.parent)
        self.INSTANCE_TIME_TAG = datetime.now().strftime('%Y%m%dT%H%M%S')
        self._run_started_perf = time.perf_counter()
        self.run_started_utc = datetime.now(timezone.utc).isoformat()
        self.scene_ledger = []
        self.timing_manifest = {
            'started_utc': self.run_started_utc,
            'execution_mode': 'serial',
            'configured_parallel': self.configured_parallel,
            'software_versions': self._runtime_versions(),
            'stages_s': self._timing_schema(),
            'scenes': {},
        }


    @property
    def input_folder(self):
        return self.settings.get('client_folder', 'inputs')['inputs']

    @property
    def output_folder(self):
        return self.settings.get('client_folder', 'output')['output']
    
    @property
    def wmask_folder(self):
        return self.settings.get('client_folder', 'wmask_folder')['wmask_folder']
    
    @property
    def roi_vectors(self):
        return self.settings['roi_vectors']
    
    @property
    def compute_l2b(self):
        return self._as_bool(
            self.settings.get('processing', {}).get('compute_l2b', False)
        )
        
    @property
    def make_report(self):
        return self._as_bool(
            self.settings.get('processing', {}).get('make_report', False)
        )
    
    @property
    def report_rrs(self):
        return self._as_bool(
            self.settings.get('processing', {}).get('report_rrs', False)
            )

    @property
    def tile_id(self):
        return self.settings.get('processing', 's2_tile')['s2_tile']

    @property
    def ac_processor(self):
        processor = self.settings.get('processing', {}).get('ac_processor', 'GRS')
        processor = str(processor).strip().upper()
        supported = ('GRS', 'ACOLITE')
        if processor not in supported:
            raise ValueError(
                f"Unsupported ac_processor {processor!r}. Supported batch processors: "
                f"{', '.join(supported)}."
            )
        return processor
    
    # @property
    # def grs_files(self):
    #     grs_file_list = u.walktalk(self.input_folder, unwanted_string='*_anc*')
    #     return grs_file_list
    
    @property
    def grs_file_version(self):
        if self.ac_processor != 'GRS':
            return None
        try:
            return self.settings['processing']['grs_version']
        except KeyError as exc:
            raise ValueError("grs_version is required when ac_processor=GRS.") from exc

    @staticmethod
    def _normalize_tile(tile):
        return str(tile).strip().upper().removeprefix('T')

    def discover_input_files(self):
        """Discover validated input products for the selected processor and tile."""
        tile = self._normalize_tile(self.tile_id)
        if self.ac_processor == 'GRS':
            search_root = os.path.join(self.input_folder, self.tile_id)
            input_files = u.walktalk(search_root, unwanted_string='*_anc*')
            metadata_reader = g.metadata
            filename_rule = '*.nc excluding *_anc*'
        else:
            search_root = self.input_folder
            filename_rule = '*_L2R.nc with ACOLITE/L2R metadata'
            candidates = sorted(Path(search_root).rglob('*_L2R.nc'))
            input_files = []
            for candidate in candidates:
                info = ACOLITE_S2.metadata(candidate, require_l2r=True)
                if self._normalize_tile(info['tile']) == tile:
                    input_files.append(candidate)
            metadata_reader = lambda path: ACOLITE_S2.metadata(path, require_l2r=True)

        records = [(path, metadata_reader(path)) for path in input_files]
        records.sort(key=lambda item: (item[1]['pydate'], str(item[0])))
        if not records:
            raise ValueError(
                f"No valid {self.ac_processor} inputs found under {search_root!r}; "
                f"expected {filename_rule} for tile {tile}."
            )
        return records

    @property
    def l2b_functions(self):
        l2b_algos = ast.literal_eval(self.settings['timeseries']['l2b_algorithms'])
        return l2b_algos

    @property
    def l2b_fx_required_bands(self):
        """
        Returns a dictionary of L2B algorithms from inversion_functions.py 
        as defined by the user in the settings.ini and their required bands.
        """
        fx_req_bands = {}
        for algo in self.l2b_functions:
            if algo not in ifunc.functions:
                raise ValueError(
                    f'Unknown L2B algorithm {algo!r}; register it in '
                    'getpak.inversion_functions.functions first.'
                )
            function = ifunc.functions[algo]['function']
            signature = inspect.signature(function)
            required_bands = [name for name, param in signature.parameters.items() if param.default == inspect.Parameter.empty]
            fx_req_bands[algo] = required_bands
        return fx_req_bands         

    @staticmethod
    def _as_bool(value):
        if isinstance(value, bool):
            return value
        if value is None:
            return False
        return str(value).strip().lower() in {'1', 'true', 'yes', 'y', 'on'}


    @property
    def mask_mode(self):
        configured_static = self._as_bool(self.settings.get('processing', {}).get('static_mask_mode', False))
        mode = str(
            self.settings.get('processing', {}).get('mask_mode', 'scene_specific')
        ).strip().lower()
        if mode not in {'scene_specific', 'static'}:
            raise ValueError("mask_mode must be 'scene_specific' or 'static'.")
        if configured_static and mode == 'scene_specific':
            return 'static'
        return mode

    @property
    def static_mask(self):
        processing = self.settings.get('processing', {})
        modern = processing.get('static_mask_path')
        legacy = processing.get('static_mask')
        if modern and legacy and Path(modern) != Path(legacy):
            raise ValueError('static_mask_path conflicts with legacy static_mask.')
        value = modern or legacy
        if self.mask_mode == 'static' and not value:
            raise ValueError('mask_mode=static requires an explicit static_mask path; static_mask_mode=True requires static_mask_path.')
        return value

    @property
    def overwrite_outputs(self):
        return self._as_bool(
            self.settings.get('processing', {}).get('overwrite_outputs', False)
        )

    @property
    def configured_parallel(self):
        return self._as_bool(
            self.settings.get('processing', {}).get('parallel', False)
        )

    @staticmethod
    def _scene_uid(processor, info):
        acquired = info['pydate']
        if getattr(acquired, 'tzinfo', None) is not None:
            acquired = acquired.astimezone(timezone.utc).replace(tzinfo=None)
        timestamp = acquired.strftime('%Y%m%dT%H%M%S')
        mission = str(info.get('mission', 'UNKNOWN')).strip().upper()
        tile = Pipelines._normalize_tile(info.get('tile', 'UNKNOWN'))
        identity = '|'.join((str(info.get('input_file')), str(info.get('product_id', '')), str(info.get('processor', processor))))
        source_tag = hashlib.sha256(identity.encode('utf-8')).hexdigest()[:16]
        return f'{processor}_{mission}_{timestamp}_T{tile}_{source_tag}'

    @staticmethod
    def _timestamp_from_name(path):
        match = re.search(
            r'(20\d{6}T\d{6})', os.path.basename(os.fspath(path))
        )
        return match.group(1) if match else None

    @staticmethod
    def _rrs_diagnostics(rrs_dict, bands):
        diagnostics = {}
        for band in bands:
            if band not in rrs_dict:
                diagnostics[band] = {
                    'neg_count': None, 'finite_count': None,
                    'neg_fraction': None, 'reason': 'missing_band',
                    'unit': 'sr-1',
                }
                continue
            values = np.asarray(rrs_dict[band].values, dtype=float)
            finite = np.isfinite(values)
            finite_count = int(np.count_nonzero(finite))
            neg_count = int(np.count_nonzero(finite & (values < 0)))
            diagnostics[band] = {
                'neg_count': neg_count,
                'finite_count': finite_count,
                'neg_fraction': (float(neg_count / finite_count)
                                 if finite_count else None),
                'reason': None if finite_count else 'no_finite_water_support',
                'unit': 'sr-1',
            }
        return diagnostics

    @staticmethod
    def _runtime_versions():
        versions = {'python': platform.python_version()}
        for distribution in (
            'getpak', 'numpy', 'pandas', 'xarray', 'dask', 'rasterio',
            "h5py", "h5netcdf",
            'rioxarray', 'rasterstats', 'GDAL', 'scikit-learn',
        ):
            try:
                versions[distribution] = package_metadata.version(distribution)
            except package_metadata.PackageNotFoundError:
                versions[distribution] = 'not-installed'
        try:
            import h5py
            versions["HDF5"] = h5py.version.hdf5_version
        except ImportError:
            versions["HDF5"] = "not-installed"
        return versions

    @staticmethod
    def _timing_schema():
        return {
            name: 0.0 for name in (
                'discovery', 'read', 'mask_reproject',
                'filtering_owt_inversion', 'raster_write',
                'roi_report', 'total',
            )
        }

    def _ledger_target(self):
        return os.path.join(
            self.output_folder, self.tile_id,
            f'{self.INSTANCE_TIME_TAG}_scene_ledger.json',
        )

    def _timing_target(self):
        return os.path.join(
            self.output_folder, self.tile_id,
            f'{self.INSTANCE_TIME_TAG}_timing_manifest.json',
        )

    @staticmethod
    def _write_json(path, payload):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        temporary = f'{path}.tmp'
        with open(temporary, 'w', encoding='utf-8') as target:
            json.dump(payload, target, indent=2, sort_keys=True)
        os.replace(temporary, path)

    def _write_scene_ledger(self):
        self.scene_ledger_path = self._ledger_target()
        self._write_json(self.scene_ledger_path, self.scene_ledger)

    def _write_timing_manifest(self):
        self.timing_manifest_path = self._timing_target()
        self._write_json(self.timing_manifest_path, self.timing_manifest)

    def _processor_version(self, info=None):
        info = info or {}
        value = info.get('processor_version') or info.get('grs_ver')
        if self.ac_processor == 'GRS':
            if str(value or '').strip().upper() in {'', 'NA', 'N/A', 'UNKNOWN'}:
                value = self.grs_file_version
        return str(value or 'unknown')

    @staticmethod
    def _acquisition_fields(info):
        acquired = info['pydate']
        if getattr(acquired, 'tzinfo', None) is not None:
            acquired = acquired.astimezone(timezone.utc).replace(tzinfo=None)
        return {
            'acquisition_datetime_utc': acquired.strftime('%Y-%m-%d %H:%M:%S'),
            'acquisition_date': acquired.strftime('%Y-%m-%d'),
            'acquisition_time_utc': acquired.strftime('%H:%M:%S'),
        }

    @staticmethod
    def _tag_value(value):
        return '' if value is None else str(value)

    def _output_contract_metadata(self, scene_uid, product, *, unit,
                                  stored_multiplier, nodata, categorical=False):
        entry = getattr(self, 'ledger_by_uid', {}).get(scene_uid, {})
        info = getattr(self, 'meta', {}).get(scene_uid, {})
        acquired = self._acquisition_fields(info) if info else {
            'acquisition_datetime_utc': entry.get('acquisition_datetime_utc', ''),
            'acquisition_date': entry.get('acquisition_date', ''),
            'acquisition_time_utc': entry.get('acquisition_time_utc', ''),
        }
        if not acquired.get("acquisition_datetime_utc"):
            raise ValueError(
                f"Cannot write {product} for {scene_uid}: source acquisition time is missing."
            )
        try:
            datetime.strptime(acquired["acquisition_datetime_utc"], "%Y-%m-%d %H:%M:%S")
        except ValueError as exc:
            raise ValueError(
                f"Cannot write {product} for {scene_uid}: acquisition time is invalid."
            ) from exc
        metadata = {
            'product': product,
            'product_family': 'categorical' if categorical else 'continuous',
            'physical_unit': unit,
            'stored_multiplier': float(stored_multiplier),
            'decode_multiplier': float(1.0 / stored_multiplier),
            'add_offset': 0.0,
            'nodata': int(nodata),
            'encoding_version': 'GETPAK-ENC-2',
            'encoding_profile': self.encoding_settings['output_encoding']['encoding_profile'],
            'dtype': self.encoding_settings['output_encoding'][
                'categorical_dtype' if categorical else 'continuous_dtype'
            ],
            'raster_scale': float(1.0 / stored_multiplier),
            'resolution': float(1.0 / stored_multiplier),
            'maximum_physical_value': (
                None if categorical else float(65534.0 / stored_multiplier)
            ),
            'acquisition_datetime_utc': acquired['acquisition_datetime_utc'],
            'acquisition_date': acquired['acquisition_date'],
            'acquisition_time_utc': acquired['acquisition_time_utc'],
            'tile': self._normalize_tile(entry.get('tile', self.tile_id)),
            'platform': entry.get('platform', entry.get('mission', info.get('mission', 'UNKNOWN'))),
            'processor': entry.get('processor', self.ac_processor),
            'processor_version': entry.get('processor_version', self._processor_version(info)),
            'scene_uid': entry.get('scene_uid', scene_uid),
            'record_id': entry.get('record_id', scene_uid),
            'source_product_name': entry.get('source_product_name', info.get('basename', '')),
        }
        grid_validation = getattr(self, 'grid_by_uid', {}).get(scene_uid)
        if grid_validation:
            metadata['grid_validation'] = json.dumps(
                grid_validation, sort_keys=True, separators=(',', ':')
            )
        return metadata

    def _output_filename(self, folder, prefix, scene_uid, suffix):
        entry = getattr(self, "ledger_by_uid", {}).get(scene_uid, {})
        product = str(prefix).rstrip('_')
        if product == 'OWTs':
            product = 'OWT'
        record_id = entry.get('record_id', scene_uid)
        timestamp = entry.get('acquisition_datetime_utc', '')
        timestamp = timestamp.replace('-', '').replace(':', '').replace(' ', 'T')[:15]
        if not timestamp:
            raise ValueError(
                f"Cannot generate {product} filename for {scene_uid}: source acquisition time is missing."
            )
        tile = self._normalize_tile(entry.get('tile', self.tile_id))
        return os.path.join(
            self.output_folder, self.tile_id, folder,
            f'{product}_{timestamp}_T{tile}_{record_id}{suffix}',
        )

    def _existing_output_identity(self, target):
        try:
            import rasterio
            with rasterio.open(target) as source:
                tags = source.tags()
        except Exception as exc:
            raise FileExistsError(
                f'Cannot establish collision identity for existing output: {target}'
            ) from exc
        required = {
            'SCENE_UID': tags.get('SCENE_UID'),
            'RECORD_ID': tags.get('RECORD_ID'),
            'PROCESSOR': tags.get('PROCESSOR'),
            'PROCESSOR_VERSION': tags.get('PROCESSOR_VERSION'),
            'GETPAK_ENCODING_VERSION': tags.get('GETPAK_ENCODING_VERSION'),
        }
        if any(value in (None, '') for value in required.values()):
            raise FileExistsError(
                f'Cannot establish required collision identity for existing output: {target}'
            )
        return required

    def _output_target(self, folder, prefix, scene_uid, suffix):
        target = self._output_filename(folder, prefix, scene_uid, suffix)
        if not os.path.exists(target):
            return target

        # Text sidecars have deterministic scene/record targets but no raster tags.
        if suffix.lower() != '.tif':
            if not self.overwrite_outputs:
                self._skip_output_targets.add(os.path.abspath(target))
            return target

        entry = getattr(self, 'ledger_by_uid', {}).get(scene_uid, {})
        expected = {
            'SCENE_UID': entry.get('scene_uid'),
            'RECORD_ID': entry.get('record_id', scene_uid),
            'PROCESSOR': entry.get('processor', self.ac_processor),
            'PROCESSOR_VERSION': entry.get('processor_version'),
            'GETPAK_ENCODING_VERSION': 'GETPAK-ENC-2',
        }
        if any(value in (None, '') for value in expected.values()):
            raise FileExistsError(
                f'Cannot establish collision identity for existing output: {target}'
            )
        actual = self._existing_output_identity(target)
        expected = {key: self._tag_value(value) for key, value in expected.items()}
        if actual != expected:
            raise FileExistsError(
                f'Refusing to overwrite conflicting output identity: {target}'
            )
        if not self.overwrite_outputs:
            self._skip_output_targets.add(os.path.abspath(target))
        return target

    def _target_was_skipped(self, target):
        return os.path.abspath(os.fspath(target)) in self._skip_output_targets

    def _write_scaled_raster(self, values, folder, prefix, scene_uid, *,
                             scale, unit, transform, projection):
        target = self._output_target(folder, prefix, scene_uid, '.tif')
        if self._target_was_skipped(target):
            return target, {'skipped_existing': True}
        nodata = self.encoding_settings['output_encoding']['continuous_nodata']
        encoded, metadata = u.to_uint16_scaled(
            values, scale=scale, nodata=nodata, unit=unit,
            product=str(prefix).rstrip('_'), return_metadata=True,
        )
        metadata.update(self._output_contract_metadata(
            scene_uid, str(prefix).rstrip('_'), unit=unit,
            stored_multiplier=scale, nodata=nodata,
        ))
        r.array2tiff(
            ndarray_data=encoded,
            str_output_file=target,
            transform=transform,
            projection=projection,
            no_data=nodata,
            metadata=metadata,
        )
        return target, metadata

    def _write_categorical_raster(self, values, folder, prefix, scene_uid, *,
                                  transform, projection):
        target = self._output_target(folder, prefix, scene_uid, '.tif')
        if self._target_was_skipped(target):
            return target
        product = 'OWT' if str(prefix).rstrip('_') == 'OWTs' else str(prefix).rstrip('_')
        nodata = self.encoding_settings['output_encoding']['categorical_nodata']
        encoded, metadata = u.to_uint8_categorical(
            values, nodata=nodata, product=product, return_metadata=True,
        )
        metadata.update(self._output_contract_metadata(
            scene_uid, product, unit='class', stored_multiplier=1,
            nodata=nodata, categorical=True,
        ))
        r.array2tiff(
            ndarray_data=encoded,
            str_output_file=target,
            transform=transform,
            projection=projection,
            no_data=nodata,
            metadata=metadata,
        )
        return target

    def run(self, compute_l2b=None, make_report=None):
        """
        Run the GET-Pak processing flow.

        If compute_l2b or make_report are None, the values are read from settings.ini.
        """
        compute_l2b = self.compute_l2b if compute_l2b is None else compute_l2b
        make_report = self.make_report if make_report is None else make_report

        if compute_l2b:
            print('Compute L2B set to True.')
            self.get_matchups()
            self.matchups_to_l2b()
        
        if make_report:
            print('Generating report.')
            self.build_report()

        return self

    def get_matchups(self, do_return=False):
        """Match inputs to masks while preserving one ledger row per input."""
        u.set_gdal_driver_path()
        sep_trace = u.repeat_to_length('-', 22)
        print(sep_trace)
        print(f'Running L2B algorithms for {self.ac_processor} with WD intersection...')
        print(f'Processing tile: {self.tile_id}')
        print(f'Input {self.ac_processor} root: {self.input_folder}')

        discovery_start = time.perf_counter()
        records = self.discover_input_files()
        tile = self._normalize_tile(self.tile_id)

        if self.mask_mode == 'static':
            static_mask = Path(self.static_mask)
            if not static_mask.is_absolute():
                static_mask = Path(self.settings['_config_dir']) / static_mask
            if not static_mask.is_file():
                raise ValueError(f'Configured static_mask does not exist: {static_mask}')
            mask_records = [{'date': None, 'path': static_mask, 'tile': tile}]
        else:
            wd_dates, wd_masks = m.get_waterdetect_masks(
                input_folder=self.wmask_folder
            )
            mask_records = [
                {
                    'date': date,
                    'path': path,
                    'tile': m._tile_from_name(path),
                    'timestamp': self._timestamp_from_name(path),
                }
                for date, path in zip(wd_dates, wd_masks)
            ]

        ledger = []
        provisional = {}
        for source_path, info in records:
            scene_uid = self._scene_uid(self.ac_processor, info)
            record_id = hashlib.sha256(
                f'{scene_uid}|{Path(source_path)}'.encode('utf-8')
            ).hexdigest()[:20]
            entry = {
                'scene_uid': scene_uid, 'record_id': record_id,
                'processor': self.ac_processor,
                'processor_version': self._processor_version(info),
                'mission': str(info.get('mission', 'UNKNOWN')),
                'platform': str(info.get('mission', 'UNKNOWN')),
                'source_product_name': str(info.get('basename', Path(source_path).name)),
                'acquisition_time': info['pydate'].isoformat(),
                **self._acquisition_fields(info),
                'tile': tile,
                'source_path': str(source_path),
                'mask_path': None,
                'mask_mode': self.mask_mode,
                'status': 'unmatched',
                'reason': 'no matching date/tile mask',
                'outputs': {},
                'error': None,
            }
            if self.mask_mode == 'static':
                candidates = mask_records
            else:
                candidates = [
                    item for item in mask_records
                    if item['date'] == info['str_date']
                    and item['tile'] in {None, tile}
                ]
                if candidates:
                    acquired = info['pydate']
                    if getattr(acquired, 'tzinfo', None) is not None:
                        acquired = acquired.astimezone(timezone.utc)
                    scene_timestamp = acquired.strftime('%Y%m%dT%H%M%S')
                    exact = [item for item in candidates if item['timestamp'] == scene_timestamp]
                    if exact:
                        candidates = exact
                        entry['mask_match_type'] = 'exact'
                    elif len(candidates) == 1:
                        entry['mask_match_type'] = 'same_day_reuse'
                    elif all(item['timestamp'] for item in candidates):
                        scene_time = datetime.strptime(scene_timestamp, '%Y%m%dT%H%M%S')
                        distances = [abs((datetime.strptime(item['timestamp'], '%Y%m%dT%H%M%S') - scene_time).total_seconds()) for item in candidates]
                        nearest = min(distances)
                        candidates = [item for item, distance in zip(candidates, distances) if distance == nearest]
                        entry['mask_match_type'] = 'nearest_same_day'
                    else:
                        entry['mask_match_type'] = 'ambiguous_same_day'

            if len(candidates) > 1:
                entry.update({
                    'status': 'ambiguous_mask',
                    'reason': 'multiple masks match this scene date/tile',
                    'candidate_masks': [str(item['path']) for item in candidates],
                })
            elif len(candidates) == 1:
                mask_path = str(candidates[0]['path'])
                entry.update({
                    'mask_path': mask_path,
                    'status': 'matched',
                    'reason': None,
                })
                entry.setdefault('mask_match_type', 'static' if self.mask_mode == 'static' else 'same_day_reuse')
                provisional.setdefault(mask_path, []).append(entry)
            ledger.append(entry)

        for entries in provisional.values():
            for entry in entries:
                entry['mask_reuse_status'] = 'reused' if len(entries) > 1 else 'single_use'

        matches = {}
        meta = {}
        for entry, (source_path, info) in zip(ledger, records):
            if entry['status'] != 'matched':
                continue
            scene_uid = entry['record_id']
            if scene_uid in matches:
                entry.update({
                    'status': 'identity_collision',
                    'reason': 'stable scene identity is not unique',
                })
                continue
            matches[scene_uid] = {
                'IMG': source_path,
                'WM': Path(entry['mask_path']),
            }
            meta[scene_uid] = info

        self.scene_ledger = ledger
        self.ledger_by_uid = {item['record_id']: item for item in ledger}
        self.matches = matches
        self.str_matches = {
            key: {'IMG': str(value['IMG']), 'WM': str(value['WM'])}
            for key, value in matches.items()
        }
        self.dates = list(matches)
        self.meta = meta
        self.timing_manifest['stages_s']['discovery'] = (
            time.perf_counter() - discovery_start
        )
        self._write_scene_ledger()

        if not matches:
            raise ValueError(
                f'No {self.ac_processor}/WaterDetect matchups found for tile {tile}; '
                f'see {self.scene_ledger_path}.'
            )

        print(f'get_matchups: {len(matches)} matched of {len(records)} discovered.')
        if do_return:
            return matches, self.str_matches, self.dates, meta
        return None
    
    def matchups_to_l2b(self):
        """
        TO-DO
        """
        report_rrs = self.report_rrs
        matches = self.matches
        str_matches = self.str_matches

        results = {}
        tot = len(self.matches)
        sep_trace = u.repeat_to_length('-', 22)
        imgs_out = os.path.join(self.output_folder, self.tile_id)
        # Creating output folder structure
        Path(os.path.join(imgs_out, "npix")).mkdir(parents=True, exist_ok=True)
        Path(os.path.join(imgs_out, "OWT")).mkdir(parents=True, exist_ok=True)
        Path(os.path.join(imgs_out, "OWTSPM")).mkdir(parents=True, exist_ok=True)
        Path(os.path.join(imgs_out, "Chla")).mkdir(parents=True, exist_ok=True)
        Path(os.path.join(imgs_out, "Turb")).mkdir(parents=True, exist_ok=True)
        Path(os.path.join(imgs_out, "HySPM")).mkdir(parents=True, exist_ok=True)
        # Rrs bands
        if report_rrs:
            Path(os.path.join(imgs_out, "Aerosol")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "Blue")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "Green")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "Red")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "RedEdge1")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "RedEdge2")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "RedEdge3")).mkdir(parents=True, exist_ok=True)
            Path(os.path.join(imgs_out, "Nir2")).mkdir(parents=True, exist_ok=True)
        
        for n, key in enumerate(matches):
            print(sep_trace)
            print(f'Processing: {n+1}/{tot} - {key}')
            
            scene_start = time.perf_counter()
            scene_timing = {
                'stages_s': self._timing_schema(),
                'scene_dimensions': None,
                'valid_water_pixels': 0,
                'roi_count': len(self.roi_vectors),
            }
            results[key] = {
                'IMG': str_matches[key]['IMG'],
                'WM': str_matches[key]['WM'],
                'scene_uid': key,
                'processor': self.ac_processor,
                'source_path': str_matches[key]['IMG'],
                'status': 'processing',
                'error': None,
                'scaling': {},
            }
            ledger_entry = self.ledger_by_uid[key]
            ledger_entry['status'] = 'processing'
            
            results[key].update({'npix': 'empty'})
            results[key].update({'OWT': 'empty'})
            results[key].update({'OWTSPM': 'empty'})   
            results[key].update({'Chla': 'empty'})
            results[key].update({'Turb': 'empty'})
            results[key].update({'HySPM': 'empty'})
            # Rrs bands
            if report_rrs:
                results[key].update({'Aerosol': 'empty'})  # 443
                results[key].update({'Blue': 'empty'})     # 490
                results[key].update({'Green': 'empty'})    # 560
                results[key].update({'Red': 'empty'})      # 665
                results[key].update({'RedEdge1': 'empty'}) # 705
                results[key].update({'RedEdge2': 'empty'}) # 740
                results[key].update({'RedEdge3': 'empty'}) # 783
                results[key].update({'Nir2': 'empty'})     # 865

            rrs_source = None
            grs = None

            def close_scene_sources():
                for dataset in (grs, rrs_source):
                    if dataset is not None:
                        try:
                            dataset.close()
                        except Exception:
                            pass

            try:
                m.reset_numerical_diagnostics()
                read_start = time.perf_counter()
                grs_ver = self.grs_file_version
                version_message = f" using grs_version={grs_ver}" if grs_ver else ""
                print(f'Loading {self.ac_processor} data{version_message}...')
                rrs_source = i.get_input_nc(
                    file=str_matches[key]['IMG'], sensor='S2MSI',
                    AC_processor=self.ac_processor, grs_version=grs_ver,
                )
                scene_timing['stages_s']['read'] = time.perf_counter() - read_start
                scene_timing['scene_dimensions'] = [
                    int(value) for value in rrs_source['Red'].shape
                ]
                grid_validation = rrs_source.attrs.get('grid_validation')
                if self.ac_processor == 'GRS':
                    if not grid_validation:
                        raise ValueError(
                            "GRS reader did not provide validated grid metadata."
                        )
                    self.grid_by_uid[key] = grid_validation
                    ledger_entry['grid_validation'] = grid_validation
                    results[key]['grid_validation'] = grid_validation

                mask_start = time.perf_counter()
                print(f'Intersecting image with water mask...')
                grs, mask_status = m.intersect_watermask(
                    rrs_dict=rrs_source,
                    water_mask_dir=str_matches[key]['WM'],
                    require_full_coverage=self.mask_mode == 'static',
                    return_status=True,
                )
                scene_timing['stages_s']['mask_reproject'] = (
                    time.perf_counter() - mask_start
                )
                if grs is None:
                    ledger_entry['status'] = mask_status
                    raise ValueError(
                        f'Water mask compatibility failure: {mask_status}.'
                    )

                # Per-band diagnostics are measured after source decoding and water masking, before QC.
                rrs_bands = ('Aerosol', 'Blue', 'Green', 'Red', 'RedEdge1', 'RedEdge2', 'RedEdge3', 'Nir2')
                rrs_diagnostics = self._rrs_diagnostics(grs, rrs_bands)
                ledger_entry['rrs_diagnostics'] = rrs_diagnostics
                results[key]['rrs_diagnostics'] = rrs_diagnostics

                #### Before using the filters, creating a matrix to store the number of pixels
                pixels = np.array([
                ['Water_pixels', '0'],
                ['Neg_Rrs_B4', '0'],
                ['Low_Rrs', '0'],
                ['OWT_1', '0'],])    
                
                # filtering bad pixels
                processing_start = time.perf_counter()
                raster_elapsed = 0.0
                print(f'Filtering bad quality pixels...')
                grs = m.filter_pixels(rrs_dict=grs, neg_rrs='Red', low_rrs=True, low_rrs_thresh=0.002, low_rrs_bands=['Aerosol', 'Blue', 'Green', 'Red', 'RedEdge1', 'RedEdge2'])

                # Clear and get Rrs
                print(f'Obtaining Rrs from Aerosol:443, Blue:490, Green:560, Red:665, RedEdge1:705, RedEdge2:740, RedEdge3:783, and Nir2:865nm...')
                aerosol = m._quick_rrs(rrs_dict=grs, bname='Aerosol')
                blue = m._quick_rrs(rrs_dict=grs, bname='Blue')
                green = m._quick_rrs(rrs_dict=grs, bname='Green')
                red = m._quick_rrs(rrs_dict=grs, bname='Red')
                rededge1 = m._quick_rrs(rrs_dict=grs, bname='RedEdge1')
                rededge2 = m._quick_rrs(rrs_dict=grs, bname='RedEdge2')
                rededge3 = m._quick_rrs(rrs_dict=grs, bname='RedEdge3')
                nir2 = m._quick_rrs(rrs_dict=grs, bname='Nir2')

                # number of pixels
                pixels[0,1] = m.npix
                pixels[1,1] = m.negpix
                pixels[2,1] = m.lowrrs

                # classifying all valid pixels
                print(f'Classifying the OWT of each pixel...')
                class_px, angles = m.classify_owt_chla_px(rrs_dict=grs, sensor='S2MSI', B1=True)
                
                if class_px.sum()<=0:
                    ledger_entry['status'] = 'no_valid_pixel'
                    raise ValueError("No valid pixels remain after masking and quality filters.")
                else:
                    
                    # classifying the OWT of each reservoir
                    print(f'Calculating the OWT weights for each pixel and writing the raster file...')
                    owt_classes, owt_weights = m.classify_owt_chla_weights(class_px=class_px, angles=angles, n=3)

                    # OWT classes for turbidity
                    classes_turb, angles_turb = m.classify_owt_spm_px(rrs_dict=grs, sensor='S2MSI', B1=True)

                    # Different classes for low Rrs pixels
                    # mask for OWT 14
                    bands = ['Aerosol', 'Blue', 'Green', 'Red', 'RedEdge1', 'RedEdge2']
                    stacked = xr.concat([grs[var] for var in bands], dim='variable')
                    max_values = stacked.fillna(-np.inf).max(dim='variable')
                    mask = max_values < 0.005
                    # first checking if there are any pixels in the mask
                    if np.where(mask.values)[0].size > 0:
                        owt_classes[0,mask] = 14
                        owt_classes[1,mask] = 0
                        owt_classes[2,mask] = 0
                        owt_weights[0,mask] = 1
                        owt_weights[1,mask] = 0
                        owt_weights[2,mask] = 0

                    # n pixels    
                    pixels[3,1] = np.count_nonzero(owt_classes[0,:,:] == 1)

                    pixels = np.vstack((pixels, np.array([
                        [f'scene_rrs_{band}_{field}', '' if values[field] is None else str(values[field])]
                        for band, values in rrs_diagnostics.items()
                        for field in ('neg_count', 'finite_count', 'neg_fraction', 'reason')
                    ], dtype=object)))
                    write_start = time.perf_counter()
                    str_output_file = self._output_target(
                        'npix', 'npixels_', key, '.txt'
                    )
                    if not self._target_was_skipped(str_output_file):
                        np.savetxt(str_output_file, pixels, fmt='%s', delimiter=';')
                    results[key]['npix'] = str_output_file

                    str_output_file = self._write_categorical_raster(
                        np.where(np.isfinite(np.asarray(red)), owt_classes[0, :, :], np.nan),
                        'OWT', 'OWTs_', key,
                        transform=grs.attrs['trans'],
                        projection=grs.attrs['proj'],
                    )
                    results[key]['OWT'] = str_output_file

                    str_output_file = self._write_categorical_raster(
                        np.where(np.isfinite(np.asarray(red)), classes_turb, np.nan),
                        'OWTSPM', 'OWTSPM_', key,
                        transform=grs.attrs['trans'],
                        projection=grs.attrs['proj'],
                    )
                    results[key]['OWTSPM'] = str_output_file
                    raster_elapsed += time.perf_counter() - write_start

                    # generating the chla product from these classes and weights
                    print(f'Calculating the chla for each dominant OWT and then the blended chla product...')
                    chla = m.blended_chla(rrs_dict=grs, owt_classes=owt_classes, owt_weights=owt_weights, limits=True)

                    # calculating turbidity
                    print(f'Calculating turbidity...')
                    turb = m.turb(rrs_dict=grs, class_owt_spt=classes_turb, alg='owt', limits=True)

                    # calculating SPM_S3
                    print(f'Calculating Hybrid-SPM...')
                    hyspm = m.turb(rrs_dict=grs, class_owt_spt=classes_turb, alg='Hybrid', limits=True)

                    # Preserve explicit post-QC water validity for every derived product.
                    valid_product = np.isfinite(red)
                    chla = np.where(valid_product, chla, np.nan)
                    turb = np.where(valid_product, turb, np.nan)
                    hyspm = np.where(valid_product, hyspm, np.nan)

                    # removing values for OWT1
                    chla[np.where(owt_classes[0,:,:]==1)] = np.nan
                    turb[np.where(owt_classes[0,:,:]==1)] = np.nan
                    hyspm[np.where(owt_classes[0,:,:]==1)] = np.nan
                    scene_timing['valid_water_pixels'] = int(
                        np.count_nonzero(np.isfinite(red))
                    )
                    scene_timing['stages_s']['filtering_owt_inversion'] = max(
                        0.0, time.perf_counter() - processing_start - raster_elapsed
                    )

                    write_start = time.perf_counter()
                    print(f'Parameter report_rrs set to {report_rrs}')
                    if report_rrs:
                        print('Writing Rrs rasters...')
                        rrs_products = {
                            'Aerosol': aerosol, 'Blue': blue, 'Green': green,
                            'Red': red, 'RedEdge1': rededge1,
                            'RedEdge2': rededge2, 'RedEdge3': rededge3,
                            'Nir2': nir2,
                        }
                        for product, values in rrs_products.items():
                            path, scale_metadata = self._write_scaled_raster(
                                values, product, f'{product}_', key,
                                scale=self.encoding_settings['output_encoding']['rrs_multiplier'], unit='sr-1',
                                transform=grs.attrs['trans'],
                                projection=grs.attrs['proj'],
                            )
                            results[key][product] = path
                            results[key]['scaling'][product] = scale_metadata

                    print('Writing the water-quality rasters...')
                    for product, values, unit in (
                        ('Chla', chla, 'mg m-3'),
                        ('Turb', turb, 'NTU'),
                        ('HySPM', hyspm, 'mg L-1'),
                    ):
                        path, scale_metadata = self._write_scaled_raster(
                            values, product, f'{product}_', key,
                            scale=self.encoding_settings['output_encoding'][u.product_multiplier_key(product)], unit=unit,
                            transform=grs.attrs['trans'],
                            projection=grs.attrs['proj'],
                        )
                        results[key][product] = path
                        results[key]['scaling'][product] = scale_metadata
                    scene_timing['stages_s']['raster_write'] = (
                        raster_elapsed + time.perf_counter() - write_start
                    )
               
                stacked = None
                results[key]['status'] = 'success'
                ledger_entry['status'] = 'success'
                ledger_entry['reason'] = None
                ledger_entry['outputs'] = {
                    name: value for name, value in results[key].items()
                    if isinstance(value, str) and os.path.exists(value)
                }

                t_hour, t_min, t_sec,_ = u.tac()
                print(f'Done processing: {n+1}/{tot} - {key} \nExecution time: {t_hour}h : {t_min}m : {t_sec}s')
            except Exception as e:
                print(f'Error processing {key}: {e}')
                results[key]['status'] = 'error'
                results[key]['error'] = str(e)
                if isinstance(e, FileExistsError):
                    ledger_entry['status'] = 'output_collision'
                    ledger_entry['reason'] = 'existing output identity collision'
                    ledger_entry['collision'] = str(e)
                elif ledger_entry['status'] in {'matched', 'processing'}:
                    ledger_entry['status'] = 'processing_error'
                ledger_entry['error'] = str(e)
                t_hour, t_min, t_sec,_ = u.tac()
                print(f'Execution time: {t_hour}h : {t_min}m : {t_sec}s')
            finally:
                close_scene_sources()
                if hasattr(m, '_numerical_diagnostics'):
                    numerical_diagnostics = json.loads(json.dumps(m._numerical_diagnostics))
                    ledger_entry['numerical_diagnostics'] = numerical_diagnostics
                    results[key]['numerical_diagnostics'] = numerical_diagnostics
                scene_timing['stages_s']['total'] = time.perf_counter() - scene_start
                scene_timing['status'] = ledger_entry['status']
                self.timing_manifest['scenes'][key] = scene_timing
                self._write_scene_ledger()
                self._write_timing_manifest()
        
        for stage in (
            'read', 'mask_reproject', 'filtering_owt_inversion',
            'raster_write', 'roi_report',
        ):
            self.timing_manifest['stages_s'][stage] = sum(
                scene['stages_s'][stage]
                for scene in self.timing_manifest['scenes'].values()
            )
        self.timing_manifest['stages_s']['total'] = (
            time.perf_counter() - self._run_started_perf
        )
        self.timing_manifest['ended_utc'] = datetime.now(timezone.utc).isoformat()
        self._write_scene_ledger()
        self._write_timing_manifest()

        res_file_out = os.path.join(imgs_out, self.INSTANCE_TIME_TAG + '.json')
        self._write_json(res_file_out, results)
        return results

    def run_l2b_raw(self):
        raise NotImplementedError('run_l2b_raw() is retired and unsupported; use getpak -c settings.ini run.')
        """
        Run L2B algorithm defined in the settings.ini over all GRS files
        inside the client input folder.
        """
        grs_file_list = self.grs_files
        grs_ver = self.grs_file_version
        for grs_file in grs_file_list:  #TODO: vectorize this loop
            print(f'Processing GRS using version: {grs_ver} file: {grs_file}')
            t_id = u.get_s2_tile_id(grs_file)
        
            print(f'Extracting S2-MSI band data from GRS.nc file...')
            img = grs.get_grs_dict(grs_nc_file=grs_file, grs_version=grs_ver)  #TODO: automate GRS version check
            img_base_name = os.path.basename(grs_file).split('.')[0]
            
            for algo in self.l2b_fx_required_bands.keys():
                print(f'Running {algo}...')
                required_bands = self.l2b_fx_required_bands[algo]  # ex: ['Red', 'Nir2']
                # Build a dictionary mapping each required band to the corresponding image data.
                band_data = {band: img[0][band].values for band in required_bands if band in img[0]}
                # Unpack data and run the L2B algorithm
                l2b_array = ifunc.functions[algo]['function'](**band_data)
                # Define output path/file
                output_tif = os.path.join(self.output_folder, f'{img_base_name}-{algo}.tif')
                print(f'Saving L2B array to {algo}.tif')
                r.s2_to_tiff(
                    ndarray_data=l2b_array,
                    output_img=output_tif,            
                    tile_id=t_id
                )

        print('Done.')
        pass
    
    @staticmethod
    def get_uid(fname):
        name = os.path.basename(fname)
        stem = os.path.splitext(name)[0]
        return stem[len('npixels_'):] if stem.startswith('npixels_') else stem
    
    @staticmethod
    def _search_uid(uid, path):
        result = [os.path.join(path,file) for file in os.listdir(path) if uid in file]
        return result
    
    def match_file_uid(self, out_folders_path, uid):
        products = ['npix', 'OWT', 'OWTSPM', 'Chla', 'Turb', 'HySPM']
        if self.report_rrs:
            products.extend([
                'Aerosol', 'Blue', 'Green', 'Red', 'RedEdge1',
                'RedEdge2', 'RedEdge3', 'Nir2',
            ])

        matched = {}
        missing = []
        ambiguous = []
        for product in products:
            folder = os.path.join(out_folders_path, product)
            if not os.path.isdir(folder):
                matched[product] = None
                missing.append(product)
                continue
            candidates = sorted(
                os.path.join(folder, name)
                for name in os.listdir(folder)
                if uid in name
            )
            if len(candidates) == 1:
                matched[product] = candidates[0]
            elif not candidates:
                matched[product] = None
                missing.append(product)
            else:
                matched[product] = None
                ambiguous.append(product)

        matched['missing_products'] = ';'.join(missing)
        matched['ambiguous_products'] = ';'.join(ambiguous)
        matched['report_file_status'] = (
            'ambiguous_files' if ambiguous
            else ('missing_products' if missing else 'complete')
        )
        return matched

    @staticmethod
    def _datetime_from_text(value):
        parsed = pd.to_datetime(value, errors="coerce", utc=True)
        if pd.isna(parsed):
            raise ValueError(f"Invalid acquisition datetime metadata: {value!r}.")
        return parsed.to_pydatetime().astimezone(timezone.utc).replace(tzinfo=None)

    @staticmethod
    def _filename_acquisition(path):
        match = re.match(r"^[^_]+_(20\d{6}T\d{6})_T[0-9A-Z]{5}_[^/]+\.tif$", Path(path).name)
        return Pipelines._datetime_from_text(match.group(1)) if match else None

    @staticmethod
    def _raster_provenance(path):
        if not path or not str(path).lower().endswith(".tif") or not os.path.isfile(path):
            return {}
        import rasterio
        try:
            with rasterio.open(path) as source:
                tags = source.tags()
        except Exception as exc:
            return {"provenance_error": str(exc)}
        result = {}
        for key, tag in {
            "scene_uid": "SCENE_UID", "record_id": "RECORD_ID",
            "processor": "PROCESSOR", "processor_version": "PROCESSOR_VERSION",
            "platform": "PLATFORM", "tile": "TILE",
            "encoding_version": "GETPAK_ENCODING_VERSION",
            "encoding_profile": "ENCODING_PROFILE",
        }.items():
            if tags.get(tag) not in (None, ""):
                result[key] = tags[tag]
        if tags.get("ACQUISITION_DATETIME_UTC") not in (None, ""):
            result["acquisition_datetime_utc"] = Pipelines._datetime_from_text(
                tags["ACQUISITION_DATETIME_UTC"]
            )
        return result
    def _ledger_index(self):
        indexed = {}
        root = Path(self.output_folder) / self.tile_id
        for ledger_path in sorted(root.glob("*_scene_ledger.json")):
            try:
                payload = json.loads(ledger_path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            rows = payload if isinstance(payload, list) else payload.get("scenes", [])
            for source_row in rows:
                if not isinstance(source_row, dict) or not source_row.get("record_id"):
                    continue
                record_id = str(source_row["record_id"])
                row = dict(source_row)
                row["_ledger_sources"] = list(row.get("_ledger_sources", [])) + [str(ledger_path)]
                previous = indexed.get(record_id)
                if previous:
                    conflicts = set(previous.get("_ledger_conflict_fields", []))
                    for field in ("scene_uid", "source_path", "processor", "processor_version", "tile"):
                        if previous.get(field) not in (None, "") and row.get(field) not in (None, "") and str(previous[field]) != str(row[field]):
                            conflicts.add(field)
                    row["_ledger_sources"] = list(previous.get("_ledger_sources", [])) + row["_ledger_sources"]
                    if conflicts:
                        row["_ledger_conflict_fields"] = sorted(conflicts)
                indexed[record_id] = row
        return indexed

    def line_builder(self):
        output_root = os.path.join(self.output_folder, self.tile_id)
        npix_folder = os.path.join(output_root, "npix")
        if not os.path.isdir(npix_folder):
            raise ValueError("No npix output folder is available to build a report.")
        uids = sorted(
            self.get_uid(name)
            for name in os.listdir(npix_folder)
            if name.startswith("npixels_") and name.endswith(".txt")
        )
        if not uids:
            raise ValueError("No processed scenes are available to build a report.")
        ledger = self._ledger_index()
        rows = {}
        for uid in uids:
            row = self.match_file_uid(output_root, uid)
            record_id = uid.rsplit("_", 1)[-1]
            row.update({key: value for key, value in ledger.get(record_id, {}).items() if key not in {"outputs", "error"}})
            row["record_id"] = record_id
            provenance = {}
            for product in ("Chla", "Turb", "HySPM", "OWT", "OWTSPM"):
                provenance = self._raster_provenance(row.get(product))
                if provenance:
                    break
            filename_date = next((self._filename_acquisition(row.get(product)) for product in ("Chla", "Turb", "HySPM", "OWT", "OWTSPM") if row.get(product)), None)
            acquisition = provenance.get("acquisition_datetime_utc") or filename_date
            if acquisition is None:
                row["acquisition_status"] = "error"
                row["acquisition_error"] = "No valid acquisition time in raster metadata or standardized output filename; cannot report record_id=" + record_id + "."
            else:
                row["acquisition_status"] = "success"
                row["acquisition_datetime_utc"] = acquisition.strftime("%Y-%m-%d %H:%M:%S")
                row["acquisition_date"] = acquisition.strftime("%Y-%m-%d")
                row["acquisition_time_utc"] = acquisition.strftime("%H:%M:%S")
            if provenance.get("scene_uid"):
                row["scene_uid"] = provenance["scene_uid"]
            row.setdefault("scene_uid", ledger.get(record_id, {}).get("scene_uid", ""))
            rows[uid] = row
        return rows

    @staticmethod
    def _flatten_report_row(row):
        flattened = {}
        for key, value in row.items():
            if key in ("scaling", "rrs_diagnostics", "outputs") and isinstance(value, dict):
                for child, child_value in value.items():
                    if isinstance(child_value, dict):
                        for leaf, leaf_value in child_value.items():
                            flattened[f"{key}_{child}_{leaf}"] = leaf_value
                    else:
                        flattened[f"{key}_{child}"] = child_value
            elif key not in ("IMG", "WM"):
                flattened[key] = value
        return flattened

    @staticmethod
    def build_excel(itermediary_dict, file_to_save):
        rows = [Pipelines._flatten_report_row(row) for row in itermediary_dict.values()]
        leading = ["record_id", "scene_uid", "acquisition_datetime_utc", "acquisition_date", "acquisition_time_utc"]
        invalid = [row.get("record_id", "") for row in rows if row.get("acquisition_status") != "error" and not row.get("acquisition_datetime_utc")]
        if invalid:
            raise ValueError("Cannot write consolidated report without valid acquisition dates: " + ", ".join(map(str, invalid)))
        primary = ("Chla", "Turb", "HySPM")
        rrs = ("Aerosol", "Blue", "Green", "Red", "RedEdge1", "RedEdge2", "RedEdge3", "Nir2")
        statistic_suffixes = {
            "min", "max", "mean", "count", "std", "median",
            "roi_features", "roi_features_with_data", "status",
        }
        quality = {"Water_pixels", "Neg_Rrs_B4", "Low_Rrs", "OWT_1", "npix_status"}
        numeric_quality = {"Water_pixels", "Neg_Rrs_B4", "Low_Rrs", "OWT_1"}
        unique = list(dict.fromkeys(key for row in rows for key in row))
        def is_stat(key, products):
            for product in products:
                prefix = product + "_"
                if key.startswith(prefix):
                    return key[len(prefix):] in statistic_suffixes
            return False
        water_columns = list(dict.fromkeys(
            [key for key in unique if key in quality or is_stat(key, primary)]
            + [key for key in unique if is_stat(key, rrs)]
        ))
        processing_columns = [key for key in unique if key not in set(leading + water_columns) and key != "acquisition_status"]
        preferred = ("source_path", "mask_path", "processor", "platform", "processor_version", "tile", "source_product_name", "mask_mode", "mask_match_type", "report_file_status", "missing_products", "ambiguous_products", "status", "reason", "error", "acquisition_status")
        processing_columns = [key for key in preferred if key in processing_columns or (key == "acquisition_status" and key in unique)] + [key for key in processing_columns if key not in preferred]
        rows.sort(key=lambda row: (pd.to_datetime(row.get("acquisition_datetime_utc"), errors="coerce", utc=True), str(row.get("record_id", ""))))

        def frame(columns):
            result = pd.DataFrame([{column: row.get(column) for column in leading + columns} for row in rows], columns=leading + columns)
            dt = pd.to_datetime(result["acquisition_datetime_utc"], errors="coerce", utc=True).dt.tz_localize(None)
            result["acquisition_datetime_utc"] = dt
            result["acquisition_date"] = dt.dt.date
            result["acquisition_time_utc"] = dt.dt.time
            for column in columns:
                if (column in numeric_quality or column.endswith(("_count", "_min", "_max", "_mean", "_std", "_median", "_fraction", "_multiplier", "_nodata", "_features", "_features_with_data"))):
                    result[column] = pd.to_numeric(result[column], errors="coerce")
            return result

        water = frame(water_columns)
        processing = frame(processing_columns)
        Path(file_to_save).parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(file_to_save, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss", date_format="yyyy-mm-dd") as writer:
            water.to_excel(writer, sheet_name="Water quality", index=False)
            processing.to_excel(writer, sheet_name="Processing details", index=False)
            workbook = writer.book
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.comments import Comment
            for worksheet in workbook.worksheets:
                worksheet.freeze_panes = "F2"
                worksheet.auto_filter.ref = worksheet.dimensions
                worksheet.sheet_view.showGridLines = False
                for cell in worksheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                if worksheet.title == "Water quality":
                    unit_notes = {
                        "Chla": "Chl-a statistics use physical units of mg m-3; *_count and *_roi_features are counts; *_status is text.",
                        "Turb": "Turbidity statistics use physical units of NTU; *_count and *_roi_features are counts; *_status is text.",
                        "HySPM": "HySPM statistics use physical units of mg L-1; *_count and *_roi_features are counts; *_status is text.",
                        "Aerosol": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "Blue": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "Green": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "Red": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "RedEdge1": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "RedEdge2": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "RedEdge3": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                        "Nir2": "Rrs statistics use physical units of sr-1; *_count and *_roi_features are counts; *_status is text.",
                    }
                    for cell in worksheet[1]:
                        for prefix, note in unit_notes.items():
                            if str(cell.value).startswith(prefix + "_"):
                                cell.comment = Comment(note, "GET-Pak")
                                break
                for cell in worksheet[1]:
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                worksheet.row_dimensions[1].height = 30
                for column in worksheet.iter_cols(1, worksheet.max_column):
                    name = worksheet.cell(1, column[0].column).value or ""
                    width = min(42, max(12, len(str(name)) + 2, max((len(str(cell.value)) for cell in column[1:25] if cell.value is not None), default=0) + 2))
                    worksheet.column_dimensions[column[0].column_letter].width = width
                for cell in worksheet["A"][1:] + worksheet["B"][1:]:
                    cell.number_format = "@"
                for cell in worksheet["C"][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                for cell in worksheet["D"][1:]:
                    cell.number_format = "yyyy-mm-dd"
                for cell in worksheet["E"][1:]:
                    cell.number_format = "hh:mm:ss"
            workbook.active = 0

    # ,---------,
    # | PARSERS |
    # '---------'
    
    @staticmethod
    def _parse_npix(path_to_npix):
        if not path_to_npix or not os.path.isfile(path_to_npix):
            return {'npix_status': 'missing_product'}
        try:
            df = pd.read_csv(path_to_npix, sep=';', header=None).T
            df.columns = df.iloc[0]
            df.drop(0, axis=0, inplace=True)
            values = {key: value[1] for key, value in df.to_dict().items()}
            values['npix_status'] = 'success'
            return values
        except Exception as error:
            return {'npix_status': 'error', 'npix_error': str(error)}
    
    @staticmethod
    def _parse_tifs(path_to_tif, shp_file, prefix='var', encoding_settings=None):
        empty = {
            f"{prefix}_{name}": None
            for name in ("min", "max", "mean", "count", "std", "median")
        }
        if not path_to_tif or not os.path.isfile(path_to_tif):
            empty[f"{prefix}_status"] = "missing_product"
            return empty

        try:
            import rasterio
            with rasterio.open(path_to_tif) as source:
                tags = source.tags()
                nodata = source.nodata
                reader_scale = source.scales[0] if source.scales else None
                reader_offset = source.offsets[0] if source.offsets else None

            version = tags.get("GETPAK_ENCODING_VERSION")
            stored_raw = tags.get("STORED_MULTIPLIER")
            decode_raw = tags.get("DECODE_MULTIPLIER")
            raster_raw = tags.get("RASTER_SCALE")
            offset_raw = tags.get("ADD_OFFSET")
            # Rasterio reports scale=1.0 and offset=0.0 for an untagged raster.
            # Those defaults are not authoritative GET-Pak encoding metadata.
            if raster_raw in (None, "") and reader_scale is not None:
                try:
                    if not np.isclose(float(reader_scale), 1.0):
                        raster_raw = reader_scale
                except (TypeError, ValueError):
                    raster_raw = reader_scale
            if offset_raw in (None, "") and reader_offset is not None:
                try:
                    if not np.isclose(float(reader_offset), 0.0):
                        offset_raw = reader_offset
                except (TypeError, ValueError):
                    offset_raw = reader_offset

            if raster_raw not in (None, "") and reader_scale is not None:
                try:
                    reader_scale_value = float(reader_scale)
                    tagged_scale_value = float(raster_raw)
                except (TypeError, ValueError):
                    reader_scale_value = tagged_scale_value = None
                if (reader_scale_value is not None and tagged_scale_value is not None
                        and not np.isclose(reader_scale_value, 1.0)
                        and not np.isclose(reader_scale_value, tagged_scale_value, rtol=1e-9, atol=1e-12)):
                    raise ValueError(f"{prefix} raster has conflicting explicit raster scale metadata.")
            if offset_raw not in (None, "") and reader_offset is not None:
                try:
                    reader_offset_value = float(reader_offset)
                    tagged_offset_value = float(offset_raw)
                except (TypeError, ValueError):
                    reader_offset_value = tagged_offset_value = None
                if (reader_offset_value is not None and tagged_offset_value is not None
                        and not np.isclose(reader_offset_value, 0.0)
                        and not np.isclose(reader_offset_value, tagged_offset_value, rtol=1e-9, atol=1e-12)):
                    raise ValueError(f"{prefix} raster has conflicting explicit offset metadata.")
            if version not in (None, "") and version != "GETPAK-ENC-2":
                raise ValueError(
                    f"{prefix} raster has invalid or unsupported embedded encoding version: {version!r}."
                )
            factor_fields = (
                ("STORED_MULTIPLIER", stored_raw),
                ("DECODE_MULTIPLIER", decode_raw),
                ("RASTER_SCALE", raster_raw),
            )
            present_factors = [(name, value) for name, value in factor_fields
                               if value not in (None, "")]
            parsed_factors = {}
            for name, value in present_factors:
                try:
                    parsed = float(value)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"{prefix} raster has an invalid {name}.") from exc
                if not np.isfinite(parsed) or parsed <= 0:
                    raise ValueError(f"{prefix} raster has an invalid {name}.")
                parsed_factors[name] = parsed

            if present_factors:
                stored_multiplier = parsed_factors.get("STORED_MULTIPLIER")
                decode_values = [
                    (name, parsed_factors[name])
                    for name in ("DECODE_MULTIPLIER", "RASTER_SCALE")
                    if name in parsed_factors
                ]
                if stored_multiplier is not None:
                    decode_multiplier = 1.0 / stored_multiplier
                elif decode_values:
                    decode_multiplier = decode_values[0][1]
                    stored_multiplier = 1.0 / decode_multiplier
                else:
                    raise ValueError(f"{prefix} raster has no usable multiplier metadata.")
                for name, value in decode_values:
                    if not np.isclose(value, decode_multiplier, rtol=1e-9, atol=1e-12):
                        raise ValueError(f"{prefix} raster has contradictory decode metadata." if name == "DECODE_MULTIPLIER" else f"{prefix} raster has contradictory raster scale metadata.")
                decode_source = "embedded_metadata"
                warning = None
            else:
                # Version, descriptive tags, and an explicit zero offset do not
                # supply a decoding factor; use the product setting visibly.
                settings = u.resolve_encoding_settings(dict(encoding_settings or {}))
                stored_multiplier = float(u.output_multiplier(settings, prefix))
                decode_multiplier = 1.0 / stored_multiplier
                decode_source = "settings_fallback"
                warning = (
                    f"{prefix} raster has no authoritative embedded multiplier; "
                    f"settings_fallback multiplier {stored_multiplier:g} was used."
                )

            if offset_raw not in (None, ""):
                try:
                    offset = float(offset_raw)
                except (TypeError, ValueError) as exc:
                    raise ValueError(
                        f"{prefix} raster has a conflicting nonzero or invalid ADD_OFFSET."
                    ) from exc
                if not np.isfinite(offset) or not np.isclose(offset, 0.0):
                    raise ValueError(
                        f"{prefix} raster has a conflicting nonzero or invalid ADD_OFFSET."
                    )

            if nodata is None:
                raise ValueError(f"{prefix} raster has no nodata value; cannot mask invalid pixels safely.")
            stats = m.shp_stats(tif_file=path_to_tif, shp_poly=shp_file)
        except Exception as error:
            empty[f"{prefix}_status"] = "error"
            empty[f"{prefix}_error"] = str(error)
            return empty

        def _decode(value, digits=6):
            return None if value is None else round(float(value) * decode_multiplier, digits)

        result = {
            f"{prefix}_min": _decode(stats.get("min")),
            f"{prefix}_max": _decode(stats.get("max")),
            f"{prefix}_mean": _decode(stats.get("mean")),
            f"{prefix}_count": stats.get("count", 0),
            f"{prefix}_std": _decode(stats.get("std")),
            f"{prefix}_median": _decode(stats.get("median")),
            f"{prefix}_status": stats.get("roi_status", "success"),
            f"{prefix}_roi_features": stats.get("roi_features"),
            f"{prefix}_roi_features_with_data": stats.get("roi_features_with_data"),
            f"{prefix}_encoding_version": version,
            f"{prefix}_encoding_profile": tags.get("ENCODING_PROFILE"),
            f"{prefix}_physical_unit": tags.get("PHYSICAL_UNIT"),
            f"{prefix}_applied_multiplier": stored_multiplier,
            f"{prefix}_decode_multiplier": decode_multiplier,
            f"{prefix}_decode_source": decode_source,
            f"{prefix}_nodata": nodata,
        }
        if warning:
            result[f"{prefix}_warning"] = warning
        return result

    def build_report(self):
        report_start = time.perf_counter()
        report_outputs = []
        
        report_rrs = self.report_rrs

        print(f'Building intermediary dictionary with the output folder : {self.output_folder}')
        itermediary_batch_dict = self.line_builder()
        
        for roi_vector in self.roi_vectors:
            print(f'Computing timeseries inside vector: {roi_vector}')
            roi_name = os.path.basename(roi_vector).split('.')[0]

            # ,--------------,
            # | CALL PARSERS |
            # '--------------'        
            
            # One-liners to fetch pixel data inside ROI in a given path of imgs
            if report_rrs:
                print('Fetching Aerosol-443nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Aerosol'], roi_vector, prefix='Aerosol', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching Blue-490nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Blue'], roi_vector, prefix='Blue', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching Green-560nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Green'], roi_vector, prefix='Green', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching Red-665nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Red'], roi_vector, prefix='Red', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching RedEdge1-705nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['RedEdge1'], roi_vector, prefix='RedEdge1', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching RedEdge2-740nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['RedEdge2'], roi_vector, prefix='RedEdge2', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching RedEdge3-783nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['RedEdge3'], roi_vector, prefix='RedEdge3', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

                print('Fetching Nir2-865nm data..')
                _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Nir2'], roi_vector, prefix='Nir2', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
                print('Done.')

            print('Fetching SPM L2B data..')
            _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['HySPM'], roi_vector, prefix='HySPM', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
            print('Done.')

            print('Fetching Turbidity L2B data..')
            _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Turb'], roi_vector, prefix='Turb', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
            print('Done.')

            print('Fetching Chl-a L2B data..')
            _ = [itermediary_batch_dict[key].update(self._parse_tifs(itermediary_batch_dict[key]['Chla'], roi_vector, prefix='Chla', encoding_settings=self.settings)) for key in itermediary_batch_dict.keys()]
            print('Done.')

            print('Fetching L2B pixel metadata..')
            _ = [itermediary_batch_dict[key].update(self._parse_npix(itermediary_batch_dict[key]['npix'])) for key in itermediary_batch_dict.keys()]
            print('Done.')

            print(f'Writing excel file at: {self.output_folder}')
            xlsx_target = os.path.join(self.output_folder, self.tile_id, self.INSTANCE_TIME_TAG + '_' + roi_name + '.xlsx')
            if os.path.exists(xlsx_target) and not self.overwrite_outputs:
                raise FileExistsError(
                    f'Refusing to overwrite existing report: {xlsx_target}'
                )
            self.build_excel(itermediary_batch_dict, file_to_save=xlsx_target)
            pass
            report_outputs.append(xlsx_target)


        report_elapsed = time.perf_counter() - report_start
        self.timing_manifest['stages_s']['roi_report'] += report_elapsed
        self.timing_manifest['stages_s']['total'] = (
            time.perf_counter() - self._run_started_perf
        )
        self.timing_manifest['ended_utc'] = datetime.now(timezone.utc).isoformat()
        self._write_timing_manifest()
        return report_outputs
def main():
    u.tic()
    u.print_logo()
    
    p = Pipelines()
    
    p.run()

    t_hour, t_min, t_sec,_ = u.tac()
    print(f'Done. \nElapsed execution time: {t_hour}h : {t_min}m : {t_sec}s')
    return 0


if __name__=='__main__':
    raise SystemExit(main())
