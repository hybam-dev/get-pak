from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import rasterio
from affine import Affine

from getpak.automation import Pipelines
from getpak.commons import Utils
from getpak.output import Raster


def test_continuous_product_multipliers_and_overflow_policy():
    for scale, value in ((10000, 0.1234), (100, 12.34), (10, 3000.0)):
        encoded, metadata = Utils.to_uint16_scaled(
            np.array([0.0, value, 65534 / scale, 65535 / scale, np.nan, -1.0]),
            scale=scale, return_metadata=True, unit='u', product='p',
        )
        assert encoded[0] == 0
        assert encoded[1] == round(value * scale)
        assert encoded[2] == 65534
        assert encoded[3] == 65535
        assert encoded[4] == encoded[5] == 65535
        assert metadata['stored_multiplier'] == scale
        assert metadata['decode_multiplier'] == pytest.approx(1 / scale)
        assert metadata['overflow_count'] == 1
        assert metadata['finite_physical_min'] == -1.0


def test_categorical_class_zero_is_valid_and_255_is_nodata():
    encoded, metadata = Utils.to_uint8_categorical(
        np.array([0, 7, np.nan, 255, -1]), return_metadata=True
    )
    assert encoded.tolist() == [0, 7, 255, 255, 255]
    assert metadata['nodata'] == 255
    assert metadata['encoding_version'] == 'GETPAK-ENC-2'


def test_geotiff_contract_tags_and_standard_fields(tmp_path):
    output = tmp_path / 'product.tif'
    encoded, metadata = Utils.to_uint16_scaled(
        np.array([[0.0, 3000.0]]), scale=10, unit='mg L-1',
        product='HySPM', return_metadata=True
    )
    metadata.update({
        'encoding_version': 'GETPAK-ENC-2', 'scene_uid': 'GRS_S2A_UID',
        'record_id': 'abc123', 'processor': 'GRS',
        'processor_version': 'v20', 'tile': '50RKU',
        'platform': 'S2A', 'source_product_name': 'source.nc',
    })
    Raster.array2tiff(encoded, output, Affine.identity(), 'EPSG:4326',
                      no_data=65535, metadata=metadata)
    with rasterio.open(output) as source:
        assert source.dtypes == ('uint16',)
        assert source.nodata == 65535
        assert source.scales == pytest.approx((0.1,))
        assert source.offsets == pytest.approx((0.0,))
        assert source.units == ('mg L-1',)
        tags = source.tags()
        assert tags['GETPAK_ENCODING_VERSION'] == 'GETPAK-ENC-2'
        assert tags['STORED_MULTIPLIER'] == '10.0'
        assert tags['SCENE_UID'] == 'GRS_S2A_UID'
        assert tags['RECORD_ID'] == 'abc123'
        assert tags['PROCESSOR_VERSION'] == 'v20'
        assert 'SCALE_FACTOR' not in tags


def test_filename_uses_acquisition_tile_and_existing_short_id(tmp_path, monkeypatch):
    config = {
        'client_folder': {'inputs': str(tmp_path), 'output': str(tmp_path / 'out'), 'wmask_folder': str(tmp_path)},
        'processing': {'s2_tile': 'T50RKU', 'ac_processor': 'GRS', 'grs_version': 'v20'},
        'timeseries': {'l2b_algorithms': '[]'}, 'roi_vectors': [],
    }
    monkeypatch.setattr('getpak.automation.u.read_config', lambda config_path=None: config)
    pipeline = Pipelines()
    uid = 'record-key'
    pipeline.ledger_by_uid = {uid: {
        'record_id': '82e157365cf2261f6ada', 'tile': '50RKU',
        'acquisition_datetime_utc': '2021-09-11 02:55:51',
    }}
    name = Path(pipeline._output_filename('Chla', 'Chla_', uid, '.tif')).name
    assert name == 'Chla_20210911T025551_T50RKU_82e157365cf2261f6ada.tif'
    assert 'TT50RKU' not in name


def test_conflicting_collision_identity_is_never_overwritten(tmp_path, monkeypatch):
    config = {
        'client_folder': {'inputs': str(tmp_path), 'output': str(tmp_path / 'out'), 'wmask_folder': str(tmp_path)},
        'processing': {'s2_tile': '50RKU', 'ac_processor': 'GRS', 'grs_version': 'v20', 'overwrite_outputs': True},
        'timeseries': {'l2b_algorithms': '[]'}, 'roi_vectors': [],
    }
    monkeypatch.setattr('getpak.automation.u.read_config', lambda config_path=None: config)
    pipeline = Pipelines()
    pipeline.ledger_by_uid = {'rid': {
        'scene_uid': 'scene-new', 'record_id': 'rid', 'processor': 'GRS',
        'processor_version': 'v20', 'tile': '50RKU',
        'acquisition_datetime_utc': '2021-09-11 02:55:51',
    }}
    target = Path(pipeline._output_filename('Chla', 'Chla_', 'rid', '.tif'))
    target.parent.mkdir(parents=True)
    with rasterio.open(target, 'w', driver='GTiff', height=1, width=1, count=1,
                       dtype='uint16', crs='EPSG:4326', transform=Affine.identity(),
                       nodata=65535) as dst:
        dst.write(np.array([[1]], dtype='uint16'), 1)
        dst.update_tags(SCENE_UID='different-scene', RECORD_ID='other', PROCESSOR='GRS',
                        PROCESSOR_VERSION='v20', GETPAK_ENCODING_VERSION='GETPAK-ENC-2')
    with pytest.raises(FileExistsError, match='conflicting output identity'):
        pipeline._output_target('Chla', 'Chla_', 'rid', '.tif')


def test_excel_report_uses_acquisition_chronology(tmp_path):
    target = tmp_path / 'report.xlsx'
    Pipelines.build_excel({
        'later': {'record_id': 'later', 'acquisition_datetime_utc': '2021-09-12 02:00:00', 'acquisition_date': '2021-09-12', 'acquisition_time_utc': '02:00:00'},
        'earlier': {'record_id': 'earlier', 'acquisition_datetime_utc': '2021-09-11 02:00:00', 'acquisition_date': '2021-09-11', 'acquisition_time_utc': '02:00:00'},
    }, target)
    frame = pd.read_excel(target)
    assert frame['record_id'].tolist() == ['earlier', 'later']
    assert pd.api.types.is_datetime64_any_dtype(frame['acquisition_datetime_utc'])
    assert pd.api.types.is_datetime64_any_dtype(frame['acquisition_date'])


def test_grs_processor_version_uses_configured_version_when_metadata_is_placeholder(tmp_path, monkeypatch):
    config = {
        'client_folder': {'inputs': str(tmp_path), 'output': str(tmp_path / 'out'), 'wmask_folder': str(tmp_path)},
        'processing': {'s2_tile': '50RKU', 'ac_processor': 'GRS', 'grs_version': 'v20'},
        'timeseries': {'l2b_algorithms': '[]'}, 'roi_vectors': [],
    }
    monkeypatch.setattr('getpak.automation.u.read_config', lambda config_path=None: config)
    pipeline = Pipelines()
    assert pipeline._processor_version({'grs_ver': 'NA'}) == 'v20'


def _full_encoding_config(**output_overrides):
    output = dict(Utils.OUTPUT_ENCODING_DEFAULTS)
    output.update(output_overrides)
    return {"output_encoding": output}


def test_encoding_defaults_are_added_when_sections_are_absent():
    settings = Utils.resolve_encoding_settings({"processing": {}})
    assert settings["output_encoding"]["encoding_profile"] == "standard"
    assert settings["output_encoding"]["turbidity_multiplier"] == 10
    assert "legacy_decoding" not in settings
    assert settings["output_encoding"]["continuous_dtype"] == "uint16"


def test_legacy_section_is_rejected_instead_of_silently_retained():
    with pytest.raises(ValueError, match="legacy_decoding.*retired"):
        Utils.resolve_encoding_settings({"legacy_decoding": {}})


def test_custom_multipliers_expose_resolution_and_maximum():
    settings = Utils.resolve_encoding_settings(_full_encoding_config(turbidity_multiplier="5", hyspm_multiplier="5"))
    output = settings["output_encoding"]
    assert output["encoding_profile"] == "custom"
    assert output["resolution"]["turbidity_multiplier"] == pytest.approx(0.2)
    assert output["maximum_physical_value"]["turbidity_multiplier"] == pytest.approx(13106.8)
    assert Utils.encoding_summary(settings)[2]["maximum"] == pytest.approx(13106.8)


@pytest.mark.parametrize(
    ("key", "value", "message"),
    (
        ("turbidity_multiplier", "0", "finite positive"),
        ("hyspm_multiplier", "-1", "finite positive"),
        ("rrs_multiplier", "nan", "finite positive"),
        ("chla_multiplier", "inf", "finite positive"),
        ("continuous_dtype", "uint32", "continuous_dtype must be uint16"),
        ("continuous_nodata", "0", "continuous_nodata must be 65535"),
        ("categorical_dtype", "uint16", "categorical_dtype must be uint8"),
        ("categorical_nodata", "0", "categorical_nodata must be 255"),
    ),
)
def test_invalid_encoding_settings_are_rejected(key, value, message):
    with pytest.raises(ValueError, match=message):
        Utils.resolve_encoding_settings(_full_encoding_config(**{key: value}))


def test_tagged_metadata_precedes_settings_fallback(tmp_path, monkeypatch):
    target = tmp_path / "tagged.tif"
    with rasterio.open(target, "w", driver="GTiff", height=1, width=1, count=1, dtype="uint16", crs="EPSG:4326", transform=Affine.identity(), nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.scales = (0.1,)
        dst.update_tags(GETPAK_ENCODING_VERSION="GETPAK-ENC-2", STORED_MULTIPLIER="10", DECODE_MULTIPLIER="0.1", RASTER_SCALE="0.1", ADD_OFFSET="0")
    monkeypatch.setattr("getpak.automation.m.shp_stats", lambda **kwargs: {"min": 100, "max": 100, "mean": 100, "count": 1, "std": 0, "median": 100})
    result = Pipelines._parse_tifs(target, "roi.shp", prefix="Turb", encoding_settings=_full_encoding_config(turbidity_multiplier=100))
    assert result["Turb_mean"] == pytest.approx(10.0)
    assert result["Turb_decode_source"] == "embedded_metadata"
    assert result["Turb_applied_multiplier"] == pytest.approx(10)


def test_untagged_default_scale_uses_settings_fallback(tmp_path, monkeypatch):
    target = tmp_path / "untagged.tif"
    with rasterio.open(target, "w", driver="GTiff", height=1, width=1, count=1, dtype="uint16", crs="EPSG:4326", transform=Affine.identity(), nodata=65535) as dst:
        dst.write(np.array([[1234]], dtype="uint16"), 1)
    with rasterio.open(target) as source:
        assert source.scales == pytest.approx((1.0,))
    monkeypatch.setattr("getpak.automation.m.shp_stats", lambda **kwargs: {"min": 1234, "max": 1234, "mean": 1234, "count": 1, "std": 0, "median": 1234})
    result = Pipelines._parse_tifs(target, "roi.shp", prefix="Turb", encoding_settings=_full_encoding_config())
    assert result["Turb_mean"] == pytest.approx(123.4)
    assert result["Turb_decode_source"] == "settings_fallback"
    assert "settings_fallback multiplier 10" in result["Turb_warning"]


def test_invalid_embedded_metadata_is_visible(tmp_path):
    target = tmp_path / "invalid.tif"
    with rasterio.open(target, "w", driver="GTiff", height=1, width=1, count=1, dtype="uint16", crs="EPSG:4326", transform=Affine.identity(), nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.update_tags(GETPAK_ENCODING_VERSION="GETPAK-ENC-2", STORED_MULTIPLIER="10", DECODE_MULTIPLIER="0.2")
    result = Pipelines._parse_tifs(target, "roi.shp", prefix="Turb", encoding_settings=_full_encoding_config())
    assert result["Turb_status"] == "error"
    assert "contradictory decode metadata" in result["Turb_error"]


@pytest.mark.parametrize("overwrite", (False, True))
def test_existing_sidecar_is_deterministic_without_raster_tags(tmp_path, monkeypatch, overwrite):
    config = {
        "client_folder": {"inputs": str(tmp_path), "output": str(tmp_path / "out"), "wmask_folder": str(tmp_path)},
        "processing": {"s2_tile": "50RKU", "ac_processor": "GRS", "grs_version": "v20", "overwrite_outputs": overwrite},
        "timeseries": {"l2b_algorithms": "[]"}, "roi_vectors": [],
    }
    monkeypatch.setattr("getpak.automation.u.read_config", lambda config_path=None: config)
    pipeline = Pipelines()
    pipeline.ledger_by_uid = {"rid": {"scene_uid": "scene-rid", "record_id": "rid", "processor": "GRS", "processor_version": "v20", "tile": "50RKU", "acquisition_datetime_utc": "2021-09-11 02:55:51"}}
    target = Path(pipeline._output_filename("npix", "npixels_", "rid", ".txt"))
    target.parent.mkdir(parents=True)
    target.write_text("existing")
    assert pipeline._output_target("npix", "npixels_", "rid", ".txt") == str(target)
    assert pipeline._target_was_skipped(target) is (not overwrite)


def test_excel_report_has_two_sheets_and_leading_columns(tmp_path):
    target = tmp_path / "report.xlsx"
    rows = {
        "later": {"record_id": "later", "scene_uid": "scene-later", "acquisition_datetime_utc": "2021-09-12 02:00:00", "Chla_mean": 4.5, "Chla_encoding_version": "GETPAK-ENC-2", "Chla_physical_unit": "mg m-3", "Chla_status": "success", "Water_pixels": "12", "Chla_applied_multiplier": 100, "Chla_decode_source": "embedded_metadata", "source_path": "/input/later.nc", "rrs_diagnostics": {"Red": {"neg_count": 0}}},
        "earlier": {"record_id": "earlier", "scene_uid": "scene-earlier", "acquisition_datetime_utc": "2021-09-11 02:00:00", "Chla_mean": 0.0, "Chla_encoding_version": "GETPAK-ENC-2", "Chla_physical_unit": "mg m-3", "Chla_status": "success", "Water_pixels": "10", "npix_status": "success", "Chla_applied_multiplier": 100, "Chla_decode_source": "settings_fallback", "source_path": "/input/earlier.nc"},
    }
    Pipelines.build_excel(rows, target)
    import openpyxl
    workbook = openpyxl.load_workbook(target)
    assert workbook.sheetnames == ["Water quality", "Processing details"]
    assert workbook.active.title == "Water quality"
    assert workbook["Water quality"].freeze_panes == "F2"
    assert workbook["Processing details"].freeze_panes == "F2"
    assert workbook["Water quality"].auto_filter.ref
    leading = ["record_id", "scene_uid", "acquisition_datetime_utc", "acquisition_date", "acquisition_time_utc"]
    assert [cell.value for cell in workbook["Water quality"][1]][:5] == leading
    assert [cell.value for cell in workbook["Processing details"][1]][:5] == leading
    assert workbook["Water quality"]["A2"].value == "earlier"
    assert workbook["Water quality"]["F2"].value == 0
    assert workbook["Processing details"]["A2"].value == "earlier"
    assert workbook["Processing details"]["F1"].value == "source_path"
    water_headers = [cell.value for cell in workbook["Water quality"][1]]
    processing_headers = [cell.value for cell in workbook["Processing details"][1]]
    assert "Chla_encoding_version" not in water_headers
    assert "Chla_physical_unit" not in water_headers
    assert "Chla_applied_multiplier" not in water_headers
    assert "Chla_decode_source" not in water_headers
    assert "Chla_encoding_version" in processing_headers
    assert "Chla_physical_unit" in processing_headers
    assert "Chla_applied_multiplier" in processing_headers
    assert "Chla_decode_source" in processing_headers
    assert workbook["Processing details"].cell(2, processing_headers.index("Chla_encoding_version") + 1).value == "GETPAK-ENC-2"
    assert workbook["Processing details"].cell(2, processing_headers.index("Chla_physical_unit") + 1).value == "mg m-3"
    assert workbook["Processing details"].cell(2, processing_headers.index("Chla_applied_multiplier") + 1).value == 100
    assert workbook["Processing details"].cell(2, processing_headers.index("Chla_decode_source") + 1).value == "settings_fallback"
    assert workbook["Water quality"].cell(2, water_headers.index("npix_status") + 1).value == "success"
    assert isinstance(workbook["Water quality"].cell(2, water_headers.index("npix_status") + 1).value, str)


def test_line_builder_uses_all_ledgers_and_filename_date_fallback(tmp_path, monkeypatch):
    config = {
        "client_folder": {"inputs": str(tmp_path), "output": str(tmp_path / "out"), "wmask_folder": str(tmp_path)},
        "processing": {"s2_tile": "50RKU", "ac_processor": "GRS", "grs_version": "v20"},
        "timeseries": {"l2b_algorithms": "[]"}, "roi_vectors": [],
    }
    monkeypatch.setattr("getpak.automation.u.read_config", lambda config_path=None: config)
    pipeline = Pipelines()
    root = Path(config["client_folder"]["output"]) / "50RKU"
    (root / "npix").mkdir(parents=True)
    (root / "Chla").mkdir()
    record_id = "abc123"
    uid = "20210911T025551_T50RKU_" + record_id
    (root / "npix" / ("npixels_" + uid + ".txt")).write_text("Water_pixels;10\n")
    with rasterio.open(root / "Chla" / ("Chla_" + uid + ".tif"), "w", driver="GTiff", height=1, width=1, count=1, dtype="uint16", crs="EPSG:4326", transform=Affine.identity(), nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
    pipeline._write_json(root / "20210911T025551_scene_ledger.json", [{"record_id": record_id, "scene_uid": "old-scene", "acquisition_datetime_utc": "2021-09-11 02:55:51"}])
    pipeline._write_json(root / "20260101T000000_scene_ledger.json", [{"record_id": "unrelated", "scene_uid": "new-scene", "acquisition_datetime_utc": "2026-01-01 00:00:00"}])
    rows = pipeline.line_builder()
    assert rows[uid]["acquisition_datetime_utc"] == "2021-09-11 02:55:51"
    assert rows[uid]["acquisition_status"] == "success"
    assert rows[uid]["scene_uid"] == "old-scene"


def test_generated_output_requires_acquisition_time(tmp_path, monkeypatch):
    config = {
        "client_folder": {"inputs": str(tmp_path), "output": str(tmp_path / "out"), "wmask_folder": str(tmp_path)},
        "processing": {"s2_tile": "50RKU", "ac_processor": "GRS", "grs_version": "v20"},
        "timeseries": {"l2b_algorithms": "[]"}, "roi_vectors": [],
    }
    monkeypatch.setattr("getpak.automation.u.read_config", lambda config_path=None: config)
    pipeline = Pipelines()
    pipeline.ledger_by_uid = {"rid": {"record_id": "rid", "tile": "50RKU"}}
    with pytest.raises(ValueError, match="acquisition time is missing"):
        pipeline._output_filename("Chla", "Chla_", "rid", ".tif")


def test_current_version_and_zero_offset_without_factor_use_settings_fallback(tmp_path, monkeypatch):
    target = tmp_path / "descriptive_only.tif"
    with rasterio.open(target, "w", driver="GTiff", height=1, width=1, count=1,
                       dtype="uint16", crs="EPSG:4326", transform=Affine.identity(),
                       nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.update_tags(GETPAK_ENCODING_VERSION="GETPAK-ENC-2",
                        ENCODING_PROFILE="standard", PHYSICAL_UNIT="NTU",
                        ADD_OFFSET="0")
    monkeypatch.setattr("getpak.automation.m.shp_stats",
                        lambda **kwargs: {"min": 100, "max": 100, "mean": 100,
                                           "count": 1, "std": 0, "median": 100})
    result = Pipelines._parse_tifs(
        target, "roi.shp", prefix="Turb",
        encoding_settings=_full_encoding_config())
    assert result["Turb_mean"] == pytest.approx(10.0)
    assert result["Turb_decode_source"] == "settings_fallback"
    assert result["Turb_applied_multiplier"] == pytest.approx(10)
    assert result["Turb_encoding_version"] == "GETPAK-ENC-2"


def test_invalid_factor_and_nonzero_offset_are_errors(tmp_path):
    invalid = tmp_path / "invalid_factor.tif"
    with rasterio.open(invalid, "w", driver="GTiff", height=1, width=1, count=1,
                       dtype="uint16", crs="EPSG:4326", transform=Affine.identity(),
                       nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.update_tags(GETPAK_ENCODING_VERSION="GETPAK-ENC-2",
                        STORED_MULTIPLIER="not-a-number")
    parsed = Pipelines._parse_tifs(invalid, "roi.shp", prefix="Turb",
                                   encoding_settings=_full_encoding_config())
    assert parsed["Turb_status"] == "error"
    assert "invalid STORED_MULTIPLIER" in parsed["Turb_error"]

    offset = tmp_path / "nonzero_offset.tif"
    with rasterio.open(offset, "w", driver="GTiff", height=1, width=1, count=1,
                       dtype="uint16", crs="EPSG:4326", transform=Affine.identity(),
                       nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.update_tags(ADD_OFFSET="1")
    parsed = Pipelines._parse_tifs(offset, "roi.shp", prefix="Turb",
                                   encoding_settings=_full_encoding_config())
    assert parsed["Turb_status"] == "error"
    assert "nonzero or invalid ADD_OFFSET" in parsed["Turb_error"]


def test_explicit_tiff_scale_and_offset_conflicts_are_visible(tmp_path):
    scale_target = tmp_path / "conflicting_scale.tif"
    with rasterio.open(scale_target, "w", driver="GTiff", height=1, width=1, count=1,
                       dtype="uint16", crs="EPSG:4326", transform=Affine.identity(),
                       nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.scales = (0.1,)
        dst.update_tags(RASTER_SCALE="0.2")
    parsed = Pipelines._parse_tifs(scale_target, "roi.shp", prefix="Turb",
                                   encoding_settings=_full_encoding_config())
    assert parsed["Turb_status"] == "error"
    assert "conflicting explicit raster scale metadata" in parsed["Turb_error"]

    offset_target = tmp_path / "conflicting_offset.tif"
    with rasterio.open(offset_target, "w", driver="GTiff", height=1, width=1, count=1,
                       dtype="uint16", crs="EPSG:4326", transform=Affine.identity(),
                       nodata=65535) as dst:
        dst.write(np.array([[100]], dtype="uint16"), 1)
        dst.offsets = (1.0,)
        dst.update_tags(ADD_OFFSET="0")
    parsed = Pipelines._parse_tifs(offset_target, "roi.shp", prefix="Turb",
                                   encoding_settings=_full_encoding_config())
    assert parsed["Turb_status"] == "error"
    assert "conflicting explicit offset metadata" in parsed["Turb_error"]
